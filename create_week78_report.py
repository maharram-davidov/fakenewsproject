"""
create_week78_report.py
======================
Generates a Week 7–8 progress report (.docx), matching the Week 5 report style.

Requires: python-docx (see requirements.txt)

Run:
    python create_week78_report.py
    python create_week78_report.py --output reports/Week7_8_Report.docx --team "Team Name"
"""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO_URL = "https://github.com/maharram-davidov/fakenewsproject"


def _set_doc_defaults(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _add_title_block(doc: Document, title: str, subtitle_lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    for line in subtitle_lines:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run(line)
    doc.add_paragraph()


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def _mono_block(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            table.rows[ri + 1].cells[ci].text = cell


def build_document(*, report_date: date, author_line: str) -> Document:
    doc = Document()
    _set_doc_defaults(doc)

    subtitle_lines = [
        "Capstone Design Project — Week 7–8 Progress Report",
        "Course: Capstone Design Project",
        "Project: AI-Based Fake News Detection System",
        f"Repository: {REPO_URL}",
        "Week: 7–8 of 12 — Telegram Bot Integration",
        report_date.strftime("Date: %B %d, %Y"),
        author_line,
    ]
    _add_title_block(doc, "AI-Based Fake News Detection System", subtitle_lines)

    _heading(doc, "1. Overview")
    _p(
        doc,
        "This report documents progress during Weeks 7–8 of the AI-Based Fake News "
        "Detection System capstone. Following the Streamlit web application (Weeks "
        "3–6), this milestone adds a Telegram bot interface so users can classify "
        "news articles from mobile or desktop chat without opening a browser. "
        "The bot reuses the same trained Logistic Regression model and TF-IDF "
        "vectoriser as the CLI and Streamlit apps via src.train_model.predict().",
    )

    _heading(doc, "2. Week 7–8 Deliverables")
    _bullets(
        doc,
        [
            "app/telegram_bot.py — asynchronous Telegram bot using python-telegram-bot v21",
            "Commands: /start, /help, /check, /url, /history, /stats",
            "SQLite persistence — predictions keyed by chat_id (default: data/telegram_bot.sqlite3)",
            "Polling for local dev; HTTPS webhook mode for cloud hosts (Railway / Render)",
            "requirements.txt — python-telegram-bot, python-docx",
            ".gitignore — data/*.sqlite3 excluded from Git",
            "README.md — roadmap and bot instructions updated",
        ],
    )

    _heading(doc, "3. Telegram Bot Overview")
    _heading(doc, "3.1 Runtime stack", level=2)
    _p(
        doc,
        "The bot uses python-telegram-bot’s async Application API with CommandHandler "
        "for slash commands. Model and vectoriser load once at startup (same discovery "
        "as web_app.py). If models/ is empty, the process exits after instructing "
        "python main.py.",
    )
    _heading(doc, "3.2 Commands", level=2)
    _numbered(
        doc,
        [
            "/start — Welcome and active model name.",
            "/help — Command reference.",
            "/check — Inline text or reply to a message containing the article.",
            "/url — Fetch a news URL; BeautifulSoup extracts visible text.",
            "/history — Last ten predictions for this chat.",
            "/stats — Fake vs real counts for this chat.",
        ],
    )
    _heading(doc, "3.3 Environment variables", level=2)
    _mono_block(
        doc,
        "TELEGRAM_BOT_TOKEN       (required)\n"
        "TELEGRAM_DB_PATH         (optional)\n"
        "TELEGRAM_WEBHOOK_URL     (optional)\n"
        "TELEGRAM_WEBHOOK_PATH    (optional, default tg-webhook)\n"
        "PORT                     (optional, webhook listen)",
    )

    _heading(doc, "4. Technical Architecture")
    _p(doc, "Thin wrapper over src/train_model.predict(); SQLite stores per-chat rows.")
    _mono_block(
        doc,
        "Telegram → CommandHandler → text or fetch_url_text → predict → db_insert → reply",
    )

    _heading(doc, "5. Key Design Decisions")
    _add_table(
        doc,
        ["Decision", "Rationale"],
        [
            ["python-telegram-bot async", "Webhook + polling support; active maintenance."],
            ["SQLite", "Simple persistence for capstone scope."],
            ["Plain-text replies", "Avoid parse errors from article punctuation."],
        ],
    )

    _heading(doc, "6. How to Run")
    _numbered(
        doc,
        [
            "python main.py — if models missing",
            "pip install -r requirements.txt",
            "export TELEGRAM_BOT_TOKEN=…",
            "python -m app.telegram_bot",
        ],
    )

    _heading(doc, "7. Model Benchmark (Reference)")
    _add_table(
        doc,
        ["Model", "Accuracy", "F1"],
        [
            ["Logistic Regression (active)", "99.22%", "99.25%"],
            ["Naive Bayes", "95.14%", "95.35%"],
        ],
    )

    _heading(doc, "8. Known Limitations")
    _add_table(
        doc,
        ["Limitation", "Note"],
        [
            ["English / domain bias", "Training corpus; Gemini planned later."],
            ["Ephemeral cloud disk", "SQLite may reset on redeploy."],
        ],
    )

    _heading(doc, "9. Roadmap Status")
    _add_table(
        doc,
        ["Phase", "Status"],
        [
            ["Week 1–6 ML + Streamlit", "Complete"],
            ["Week 7–8 Telegram + SQLite", "Complete"],
            ["Week 9–12 Gemini / Vision / API", "Planned"],
        ],
    )

    _heading(doc, "10. Next Steps")
    _bullets(
        doc,
        [
            "Deploy webhook with HTTPS if production Telegram traffic is required.",
            "Week 9–10: Gemini explainability and summaries.",
        ],
    )
    doc.add_paragraph()
    _p(doc, "— End of Week 7–8 Report —")
    return doc


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Week 7–8 capstone DOCX report.")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("Week7_8_Telegram_Bot_Progress_Report.docx"),
    )
    parser.add_argument("--team", default="Team — Capstone Design Project")
    parser.add_argument("--date", default="", help="YYYY-MM-DD (default: today)")
    args = parser.parse_args()

    if args.date:
        y, m, d = map(int, args.date.split("-"))
        report_date = date(y, m, d)
    else:
        report_date = date.today()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document(report_date=report_date, author_line=args.team)
    out = args.output.resolve()
    doc.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
