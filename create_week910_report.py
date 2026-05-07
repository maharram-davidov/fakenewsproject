"""
create_week910_report.py
=========================
Generates a Week 9–10 progress report (.docx): Google Gemini integration for
explainability, summaries, and Q&A (Streamlit + optional Telegram /explain).

Requires: python-docx (see requirements.txt)

Run:
    python create_week910_report.py
    python create_week910_report.py -o reports/Week910_Report.docx --team "Team Name"
"""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO_URL = "https://github.com/maharram-davidov/fakenewsproject"


def _set_doc_defaults(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _add_title_block(doc: Document, title: str, subtitle_lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    for line in subtitle_lines:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run(line)
    doc.add_paragraph()


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def _mono_block(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            table.rows[ri + 1].cells[ci].text = cell


def build_document(*, report_date: date, author_line: str) -> Document:
    doc = Document()
    _set_doc_defaults(doc)

    subtitle_lines = [
        "Capstone Design Project — Week 9–10 Progress Report",
        "Course: Capstone Design Project",
        "Project: AI-Based Fake News Detection System",
        f"Repository: {REPO_URL}",
        "Week: 9–10 of 12 — Google Gemini (explainability & chat)",
        report_date.strftime("Date: %B %d, %Y"),
        author_line,
    ]
    _add_title_block(doc, "AI-Based Fake News Detection System", subtitle_lines)

    _heading(doc, "1. Overview")
    _p(
        doc,
        "Weeks 9–10 add a Google Gemini large-language-model layer on top of the "
        "existing TF-IDF + Logistic Regression classifier. The classical model "
        "still produces the quantitative Fake/Real label and confidence; Gemini "
        "provides qualitative explanations, short summaries, and conversational "
        "Q&A grounded in the article text. This aligns with the capstone roadmap: "
        "human-readable rationale alongside statistical NLP — without replacing "
        "ground-truth fact-checking.",
    )

    _heading(doc, "2. Week 9–10 Deliverables")
    _bullets(
        doc,
        [
            "src/gemini_helper.py — prompt builders (explain / summarise / chat), truncation, generate_gemini_text()",
            "app/web_app.py — new Streamlit tab “Gemini AI”: Explain verdict, Summarise, chat UI",
            "app/telegram_bot.py — optional /explain command when GEMINI_API_KEY is set on the server",
            "requirements.txt — google-generativeai + protobuf pin compatible with Streamlit",
            "README.md — GEMINI_API_KEY / secrets.toml documentation",
        ],
    )

    _heading(doc, "3. Gemini Integration Overview")

    _heading(doc, "3.1 Authentication & configuration", level=2)
    _p(
        doc,
        "The API key is read from the environment (GEMINI_API_KEY or GOOGLE_API_KEY) "
        "or from Streamlit secrets (.streamlit/secrets.toml). Keys must never be "
        "committed to Git.",
    )
    _mono_block(
        doc,
        "GEMINI_API_KEY or GOOGLE_API_KEY\n"
        "GEMINI_MODEL (optional; else tries gemini-2.0-flash, gemini-1.5-flash)\n"
        "GEMINI_MAX_CHARS (optional; truncates long articles)",
    )

    _heading(doc, "3.2 Streamlit — Gemini AI tab", level=2)
    _numbered(
        doc,
        [
            "Article textarea plus “Load last analysed article” from Text/URL tabs.",
            "Explain ML verdict — requires a prior prediction (last_ml_result) and article text; calls build_explain_prompt().",
            "Summarise article — headline-style takeaway + bullets via build_summarise_prompt().",
            "Ask questions — st.chat_input + gemini_msgs session history; build_chat_prompt().",
        ],
    )

    _heading(doc, "3.3 Telegram — /explain", level=2)
    _p(
        doc,
        "If the bot host exports GEMINI_API_KEY, users can reply to an article "
        "message with /explain or send /explain followed by text. The handler runs "
        "predict() then Gemini with the same explain prompt pattern as the web app.",
    )

    _heading(doc, "3.4 Prompt design principles", level=2)
    _bullets(
        doc,
        [
            "Explain — stresses ML is stylistic, not legal truth; asks for balanced cues and verification suggestions.",
            "Summarise — structured bullets + uncertainty sentence.",
            "Chat — answers must cite article scope; no fabricated quotations.",
        ],
    )

    _heading(doc, "4. Technical Architecture")
    _mono_block(
        doc,
        "Streamlit / Telegram\n"
        "    → gemini_helper.build_*_prompt(article, …)\n"
        "    → google.generativeai generate_content\n"
        "    → plain-text reply (Markdown avoided where user text could break parsing)",
    )

    _heading(doc, "5. Key Design Decisions")
    _add_table(
        doc,
        ["Decision", "Rationale"],
        [
            [
                "Gemini alongside LR, not instead",
                "Preserves reproducible baseline; Gemini adds UX value.",
            ],
            [
                "Truncate articles (~28k chars default)",
                "Cost/latency control and stable behaviour.",
            ],
            [
                "Plain-text Telegram replies",
                "Avoid Telegram Markdown failures from raw article snippets.",
            ],
            [
                "protobuf < 5 pin",
                "Keeps Streamlit 1.35 installs compatible.",
            ],
        ],
    )

    _heading(doc, "6. How to Run")
    _numbered(
        doc,
        [
            "python main.py — ensure models/*.pkl exist",
            "pip install -r requirements.txt",
            "export GEMINI_API_KEY=… (or Streamlit secrets)",
            "streamlit run app/web_app.py — open Gemini AI tab",
            "Optional: python -m app.telegram_bot with GEMINI_API_KEY for /explain",
        ],
    )

    _heading(doc, "7. Limitations & Responsible Use")
    _add_table(
        doc,
        ["Topic", "Note"],
        [
            ["Hallucination risk", "Gemini may omit nuance; users verify externally."],
            ["API quotas / cost", "Production requires budgeting and rate limits."],
            ["English bias", "Same domain limits as underlying LR training data."],
            ["Not legal verdict", "Educational assistant only."],
        ],
    )

    _heading(doc, "8. Roadmap Status")
    _add_table(
        doc,
        ["Phase", "Status"],
        [
            ["Week 1–8 ML / Streamlit / Telegram", "Complete"],
            ["Week 9–10 Gemini layer", "Complete"],
            ["Week 11–12 Vision OCR / optional BERT / REST API", "Planned"],
        ],
    )

    _heading(doc, "9. Next Steps")
    _bullets(
        doc,
        [
            "Week 11–12: Google Vision OCR for screenshots; optional REST deployment.",
            "Optional: cache Gemini responses per article hash to reduce duplicate calls.",
            "Optional: structured JSON outputs from Gemini for UI cards.",
        ],
    )

    doc.add_paragraph()
    _p(doc, "— End of Week 9–10 Report —")
    return doc


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Week 9–10 capstone DOCX report.")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("Week9_10_Gemini_Progress_Report.docx"),
        help="Output .docx path",
    )
    parser.add_argument(
        "--team",
        default="Team — Capstone Design Project",
        help="Author line on title page",
    )
    parser.add_argument("--date", default="", help="YYYY-MM-DD (default: today)")
    args = parser.parse_args()

    if args.date:
        y, m, d = map(int, args.date.split("-"))
        report_date = date(y, m, d)
    else:
        report_date = date.today()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document(report_date=report_date, author_line=args.team)
    out = args.output.resolve()
    doc.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
